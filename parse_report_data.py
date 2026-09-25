# -*- coding: utf-8 -*-
"""
Module trích xuất & tính toán dữ liệu Báo cáo Tổng hợp từ file Masterlist Excel
Hỗ trợ cả file CNTT_Masterlist 2027-2028_Mau.xlsx và Masterlist 2027-2028_Mau_Tach.xlsx
Tự động giải các công thức SUMIFS từ sheet PL1.1 nếu file có dữ liệu vật tư.
"""
import os
import re
import openpyxl

DEFAULT_CNTT_PATH = r'd:\Haint28\CONG VIEC\20260824 QHĐC Thi truong\STL\QHDC 2027-2028\Final\CNTT\CNTT_Masterlist 2027-2028_Mau.xlsx'
WORKSPACE_DIR = os.path.dirname(os.path.abspath(__file__))
FALLBACK_TACH_PATH = os.path.join(WORKSPACE_DIR, 'Masterlist 2027-2028_Mau_Tach.xlsx')
FALLBACK_MAU_PATH = os.path.join(WORKSPACE_DIR, 'Masterlist 2027-2028_Mau.xlsx')

def find_best_report_file(custom_path=None):
    if custom_path and os.path.exists(custom_path):
        return custom_path
    if os.path.exists(DEFAULT_CNTT_PATH):
        return DEFAULT_CNTT_PATH
    if os.path.exists(FALLBACK_TACH_PATH):
        return FALLBACK_TACH_PATH
    if os.path.exists(FALLBACK_MAU_PATH):
        return FALLBACK_MAU_PATH
    return None

def safe_num(val):
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    try:
        s = str(val).replace(',', '').strip()
        return float(s)
    except Exception:
        return 0.0

# Ánh xạ cột trong sheet "TH theo DV"
# 2027: D=VT, E=ML, F=CĐBR, G=CNTT, H=TD, I=IP, J=CĐ, K=HT, L=TOTAL
# 2028: N=VT, O=ML, P=CĐBR, Q=CNTT, R=TD, S=IP, T=CĐ, U=HT, V=TOTAL
MANG_COL_MAP_2027 = {
    'D': 'VT', 'E': 'ML', 'F': 'CĐBR', 'G': 'CNTT', 'H': 'TD', 'I': 'IP', 'J': 'CĐ', 'K': 'HT', 'L': 'TOTAL'
}
MANG_COL_MAP_2028 = {
    'N': 'VT', 'O': 'ML', 'P': 'CĐBR', 'Q': 'CNTT', 'R': 'TD', 'S': 'IP', 'T': 'CĐ', 'U': 'HT', 'V': 'TOTAL'
}

def norm_mang(m):
    if not m: return ''
    s = str(m).strip().upper()
    if s in ['CDBR', 'CĐBR', 'CO DINH', 'CỐ ĐỊNH', 'BRCĐ']: return 'CĐBR'
    if s in ['CD', 'CĐ', 'CO DIEN', 'CƠ ĐIỆN']: return 'CĐ'
    if s in ['VT', 'VO TUYEN', 'VÔ TUYẾN']: return 'VT'
    if s in ['ML', 'MANG LOI', 'MẠNG LÕI']: return 'ML'
    if s in ['CNTT', 'CONG NGHE THONG TIN']: return 'CNTT'
    if s in ['TD', 'TRUYEN DAN', 'TRUYỀN DẪN']: return 'TD'
    if s in ['IP']: return 'IP'
    if s in ['HT', 'HA TANG', 'HẠ TẦNG']: return 'HT'
    return s

def norm_madv(d):
    if not d: return ''
    return str(d).strip().upper().replace('Ư', 'U').replace('Đ', 'D')

def build_th_dv_matrix(wb_val):
    """Tính toán ma trận TH theo DV trực tiếp từ PL1.1 nếu có dữ liệu"""
    matrix_27 = {}
    matrix_28 = {}

    if 'PL1.1 ML2027-2028' not in wb_val.sheetnames:
        return matrix_27, matrix_28

    ws_pl = wb_val['PL1.1 ML2027-2028']
    if ws_pl.max_row <= 7:
        return matrix_27, matrix_28

    for r in range(7, ws_pl.max_row + 1):
        tt27 = ws_pl.cell(r, 7).value
        tt28 = ws_pl.cell(r, 8).value
        mang = ws_pl.cell(r, 10).value
        madv = ws_pl.cell(r, 11).value

        if not mang or (tt27 is None and tt28 is None):
            continue

        mang_key = norm_mang(mang)
        madv_key = norm_madv(madv)

        val27 = safe_num(tt27) / 1000000.0
        val28 = safe_num(tt28) / 1000000.0

        k = (mang_key, madv_key)
        matrix_27[k] = matrix_27.get(k, 0.0) + val27
        matrix_28[k] = matrix_28.get(k, 0.0) + val28

    return matrix_27, matrix_28

def parse_report_workbook(file_path):
    if not file_path or not os.path.exists(file_path):
        raise FileNotFoundError(f"Không tìm thấy file: {file_path}")

    # 1. Đọc dữ liệu vật tư và build ma trận SUMIFS từ file đầu vào của người dùng
    wb_input = openpyxl.load_workbook(file_path, data_only=True)
    matrix_27, matrix_28 = build_th_dv_matrix(wb_input)

    # 2. Kiểm tra xem file có sẵn sheet "TH theo Mảng" và "TH theo Dịch vụ" không
    wb_val = wb_input
    wb_form = openpyxl.load_workbook(file_path, data_only=False)

    sheet_names = wb_val.sheetnames
    name_m = None
    name_dv = None
    for s in sheet_names:
        s_lower = s.lower()
        if 'mảng' in s_lower or 'mang' in s_lower:
            name_m = s
        if 'dịch vụ' in s_lower or 'dich vu' in s_lower:
            name_dv = s

    # Nếu file đầu vào chỉ là file 1 mảng (không có 2 sheet này), mượn cấu trúc bảng từ file mẫu
    layout_file = file_path
    if not name_m or not name_dv:
        fallback = FALLBACK_TACH_PATH if os.path.exists(FALLBACK_TACH_PATH) else FALLBACK_MAU_PATH
        if os.path.exists(fallback):
            wb_val = openpyxl.load_workbook(fallback, data_only=True)
            wb_form = openpyxl.load_workbook(fallback, data_only=False)
            layout_file = fallback
            name_m = 'TH theo Mảng'
            name_dv = 'TH theo Dịch vụ'

    # Đọc bản đồ Mã DV theo hàng trong sheet 'TH theo DV' (từ file layout)
    dv_row_map = {}
    target_dv_sheet = 'TH theo DV' if 'TH theo DV' in wb_val.sheetnames else ('TH theo DV' if 'TH theo DV' in wb_input.sheetnames else None)
    ws_dv_ref = wb_val[target_dv_sheet] if (target_dv_sheet and target_dv_sheet in wb_val.sheetnames) else (wb_input[target_dv_sheet] if (target_dv_sheet and target_dv_sheet in wb_input.sheetnames) else None)
    if ws_dv_ref:
        for r in range(4, ws_dv_ref.max_row + 1):
            c_val = ws_dv_ref.cell(r, 3).value
            if c_val:
                dv_row_map[r] = norm_madv(c_val)

    tables_mang = parse_sheet_mang(wb_val, wb_form, name_m, matrix_27, matrix_28, dv_row_map) if name_m else []
    tables_dv = parse_sheet_dv(wb_val, wb_form, name_dv, matrix_27, matrix_28, dv_row_map) if name_dv else []
    summary_data = compute_overall_summary(wb_val, tables_mang, tables_dv)
    strategy_comp = build_strategy_comparison(tables_mang, tables_dv)
    th_theo_dv = compute_th_theo_dv(matrix_27, matrix_28)

    return {
        'success': True,
        'file_path': file_path,
        'file_name': os.path.basename(file_path),
        'file_size': os.path.getsize(file_path),
        'sheets': wb_val.sheetnames,
        'summary': summary_data,
        'tables_mang': tables_mang,
        'tables_dv': tables_dv,
        'th_theo_dv': th_theo_dv,
        'strategy_comparison': strategy_comp
    }

def resolve_cell_formula(form_val, val_val, matrix_27, matrix_28, dv_row_map):
    """Giải giá trị công thức tham chiếu 'TH theo DV'!<Col><Row> nếu có matrix từ PL1.1"""
    if form_val and isinstance(form_val, str) and (matrix_27 or matrix_28):
        matches = re.findall(r"'?TH theo DV'?!([A-Z]+)(\d+)", form_val, re.IGNORECASE)
        if matches:
            tot = 0.0
            found = False
            for col_letter, row_idx_str in matches:
                col = col_letter.upper()
                r_idx = int(row_idx_str)
                ma_dv = dv_row_map.get(r_idx, '')
                if col in MANG_COL_MAP_2027:
                    mang = MANG_COL_MAP_2027[col]
                    if mang == 'TOTAL':
                        tot += sum(v for (m, dv), v in matrix_27.items() if dv == ma_dv)
                    else:
                        tot += matrix_27.get((mang, ma_dv), 0.0)
                    found = True
                elif col in MANG_COL_MAP_2028:
                    mang = MANG_COL_MAP_2028[col]
                    if mang == 'TOTAL':
                        tot += sum(v for (m, dv), v in matrix_28.items() if dv == ma_dv)
                    else:
                        tot += matrix_28.get((mang, ma_dv), 0.0)
                    found = True
            if found:
                return tot

    if val_val is not None and isinstance(val_val, (int, float)) and val_val > 0:
        return float(val_val)

    return safe_num(val_val)

def parse_sheet_mang(wb_val, wb_form, sheet_name, matrix_27, matrix_28, dv_row_map):
    if sheet_name not in wb_val.sheetnames:
        return []

    ws_val = wb_val[sheet_name]
    ws_form = wb_form[sheet_name]

    tables = []
    current_table = None

    for r in range(1, ws_val.max_row + 1):
        b_val = ws_val.cell(r, 2).value
        c_val = ws_val.cell(r, 3).value
        d_val = ws_val.cell(r, 4).value
        e_val = ws_val.cell(r, 5).value
        f_val = ws_val.cell(r, 6).value

        b_form = ws_form.cell(r, 2).value
        c_form = ws_form.cell(r, 3).value
        d_form = ws_form.cell(r, 4).value
        e_form = ws_form.cell(r, 5).value
        f_form = ws_form.cell(r, 6).value

        b_str = str(b_val or '').strip()
        c_str = str(c_val or '').strip()

        is_sec = False
        for prefix in ['I.', 'II.', 'III.', 'IV.', 'V.', 'VI.', 'VII.', 'VIII.', 'IX.', 'X.', 'Đầu tư', 'Hiện đại']:
            if b_str.startswith(prefix) and not (b_str == 'STT' or c_str == 'STT' or c_str == 'Danh mục'):
                is_sec = True
                break

        if is_sec:
            current_table = {
                'id': f'mang_sec_{len(tables) + 1}',
                'title': b_str,
                'rows': []
            }
            tables.append(current_table)
            continue

        if b_str == 'STT' or c_str == 'STT' or c_str == 'Danh mục' or b_str.startswith('Phụ lục'):
            continue

        if current_table is not None:
            name = c_str if c_str else b_str
            if name or d_val is not None or e_val is not None or f_val is not None:
                stt = b_str if (b_str and b_str != name and b_str != c_str) else ''
                is_total = (name.lower() == 'tổng' or stt.lower() == 'tổng')
                is_sub = (stt == '-' or name.startswith('-'))

                # Giải số liệu 2027 và 2028
                y27 = resolve_cell_formula(e_form, e_val, matrix_27, matrix_28, dv_row_map)
                y28 = resolve_cell_formula(f_form, f_val, matrix_27, matrix_28, dv_row_map)
                tot = resolve_cell_formula(d_form, d_val, matrix_27, matrix_28, dv_row_map)
                if tot == 0.0 and (y27 > 0 or y28 > 0):
                    tot = y27 + y28

                strat_tot = safe_num(ws_val.cell(r, 12).value or ws_val.cell(r, 10).value)
                strat_26 = safe_num(ws_val.cell(r, 13).value or ws_val.cell(r, 11).value)
                strat_27 = safe_num(ws_val.cell(r, 14).value or ws_val.cell(r, 12).value)
                strat_28 = safe_num(ws_val.cell(r, 15).value or ws_val.cell(r, 13).value)
                strat_29 = safe_num(ws_val.cell(r, 16).value or ws_val.cell(r, 14).value)
                strat_30 = safe_num(ws_val.cell(r, 17).value or ws_val.cell(r, 15).value)

                note = str(ws_val.cell(r, 18).value or ws_val.cell(r, 16).value or ws_val.cell(r, 7).value or '').strip()

                if name and not (name.endswith(':') and tot == 0 and not is_total):
                    clean_name = name.lstrip('-').strip()
                    current_table['rows'].append({
                        'stt': stt,
                        'name': clean_name if not is_total else 'Tổng',
                        'raw_name': name,
                        'is_total': is_total,
                        'is_sub': is_sub,
                        'tong': tot,
                        'y2027': y27,
                        'y2028': y28,
                        'strat_tot': strat_tot,
                        'strat_26': strat_26,
                        'strat_27': strat_27,
                        'strat_28': strat_28,
                        'strat_29': strat_29,
                        'strat_30': strat_30,
                        'note': note
                    })

    # Cập nhật lại hàng tổng nếu các dòng con có dữ liệu
    for t in tables:
        tot_row = next((r for r in t['rows'] if r['is_total']), None)
        if tot_row and tot_row['tong'] == 0:
            sub_27 = sum(r['y2027'] for r in t['rows'] if not r['is_total'] and not r['is_sub'])
            sub_28 = sum(r['y2028'] for r in t['rows'] if not r['is_total'] and not r['is_sub'])
            if sub_27 > 0 or sub_28 > 0:
                tot_row['y2027'] = sub_27
                tot_row['y2028'] = sub_28
                tot_row['tong'] = sub_27 + sub_28

    return [t for t in tables if len(t['rows']) > 0]

def parse_sheet_dv(wb_val, wb_form, sheet_name, matrix_27, matrix_28, dv_row_map):
    if sheet_name not in wb_val.sheetnames:
        return []

    ws_val = wb_val[sheet_name]
    ws_form = wb_form[sheet_name]

    tables = []
    current_table = None

    for r in range(1, ws_val.max_row + 1):
        b_val = ws_val.cell(r, 2).value
        c_val = ws_val.cell(r, 3).value
        d_val = ws_val.cell(r, 4).value
        e_val = ws_val.cell(r, 5).value
        f_val = ws_val.cell(r, 6).value
        g_val = ws_val.cell(r, 7).value

        b_form = ws_form.cell(r, 2).value
        c_form = ws_form.cell(r, 3).value
        d_form = ws_form.cell(r, 4).value
        e_form = ws_form.cell(r, 5).value
        f_form = ws_form.cell(r, 6).value
        g_form = ws_form.cell(r, 7).value

        b_str = str(b_val or '').strip()
        c_str = str(c_val or '').strip()
        d_str = str(d_val or '').strip()

        is_sec = False
        for prefix in ['I.', 'II.', 'III.', 'IV.', 'V.', 'VI.', 'VII.', 'VIII.', 'IX.', 'X.', 'Đầu tư', 'Hiện đại']:
            if b_str.startswith(prefix) and not (b_str == 'STT' or c_str == 'STT' or c_str == 'Dịch vụ'):
                is_sec = True
                break

        if is_sec:
            current_table = {
                'id': f'dv_sec_{len(tables) + 1}',
                'title': b_str,
                'rows': []
            }
            tables.append(current_table)
            continue

        if b_str == 'STT' or c_str == 'STT' or c_str == 'Dịch vụ' or b_str.startswith('Phụ lục'):
            continue

        if current_table is not None:
            name = d_str if d_str else (c_str if c_str else b_str)
            service = c_str
            if name or e_val is not None or f_val is not None or g_val is not None:
                stt = b_str if (b_str and b_str != name and b_str != c_str and b_str != d_str) else ''
                is_total = (name.lower() == 'tổng' or d_str.lower() == 'tổng' or stt.lower() == 'tổng')
                is_sub = (stt == '-' or name.startswith('-') or d_str.startswith('-'))

                y27 = resolve_cell_formula(f_form, f_val, matrix_27, matrix_28, dv_row_map)
                y28 = resolve_cell_formula(g_form, g_val, matrix_27, matrix_28, dv_row_map)
                tot = resolve_cell_formula(e_form, e_val, matrix_27, matrix_28, dv_row_map)
                if tot == 0.0 and (y27 > 0 or y28 > 0):
                    tot = y27 + y28

                strat_tot = safe_num(ws_val.cell(r, 11).value)
                strat_26 = safe_num(ws_val.cell(r, 12).value)
                strat_27 = safe_num(ws_val.cell(r, 13).value)
                strat_28 = safe_num(ws_val.cell(r, 14).value)
                strat_29 = safe_num(ws_val.cell(r, 15).value)
                strat_30 = safe_num(ws_val.cell(r, 16).value)

                note = str(ws_val.cell(r, 17).value or ws_val.cell(r, 8).value or '').strip()

                if name and not (name.endswith(':') and tot == 0 and not is_total):
                    clean_name = name.lstrip('-').strip()
                    current_table['rows'].append({
                        'stt': stt,
                        'service': service,
                        'name': clean_name if not is_total else 'Tổng',
                        'raw_name': name,
                        'is_total': is_total,
                        'is_sub': is_sub,
                        'tong': tot,
                        'y2027': y27,
                        'y2028': y28,
                        'strat_tot': strat_tot,
                        'strat_26': strat_26,
                        'strat_27': strat_27,
                        'strat_28': strat_28,
                        'strat_29': strat_29,
                        'strat_30': strat_30,
                        'note': note
                    })

    # Cập nhật lại hàng tổng nếu các dòng con có dữ liệu
    for t in tables:
        tot_row = next((r for r in t['rows'] if r['is_total']), None)
        if tot_row and tot_row['tong'] == 0:
            sub_27 = sum(r['y2027'] for r in t['rows'] if not r['is_total'] and not r['is_sub'])
            sub_28 = sum(r['y2028'] for r in t['rows'] if not r['is_total'] and not r['is_sub'])
            if sub_27 > 0 or sub_28 > 0:
                tot_row['y2027'] = sub_27
                tot_row['y2028'] = sub_28
                tot_row['tong'] = sub_27 + sub_28

    return [t for t in tables if len(t['rows']) > 0]

def compute_overall_summary(wb_val, tables_mang, tables_dv):
    total_27 = 0.0
    total_28 = 0.0
    total_sum = 0.0
    total_strat = 0.0

    mang_summary = []
    for t in tables_mang:
        tot_row = next((r for r in t['rows'] if r['is_total']), None)
        if not tot_row and t['rows']:
            tot_val = sum(r['tong'] for r in t['rows'] if not r['is_sub'])
            y27_val = sum(r['y2027'] for r in t['rows'] if not r['is_sub'])
            y28_val = sum(r['y2028'] for r in t['rows'] if not r['is_sub'])
            strat_val = sum(r['strat_tot'] for r in t['rows'] if not r['is_sub'])
        else:
            tot_val = tot_row['tong'] if tot_row else 0.0
            y27_val = tot_row['y2027'] if tot_row else 0.0
            y28_val = tot_row['y2028'] if tot_row else 0.0
            strat_val = tot_row['strat_tot'] if tot_row else 0.0

        total_27 += y27_val
        total_28 += y28_val
        total_sum += tot_val
        total_strat += strat_val

        mang_summary.append({
            'name': t['title'],
            'y2027': y27_val,
            'y2028': y28_val,
            'tong': tot_val,
            'strat_tot': strat_val
        })

    for m in mang_summary:
        m['share'] = round((m['tong'] / total_sum * 100), 2) if total_sum > 0 else 0.0

    return {
        'total_investment': total_sum,
        'total_2027': total_27,
        'total_2028': total_28,
        'total_strategy_5y': total_strat,
        'share_2027': round((total_27 / total_sum * 100), 1) if total_sum > 0 else 0.0,
        'share_2028': round((total_28 / total_sum * 100), 1) if total_sum > 0 else 0.0,
        'by_mang': mang_summary
    }

def build_strategy_comparison(tables_mang, tables_dv):
    comparison = []
    for t in tables_mang:
        tot_row = next((r for r in t['rows'] if r['is_total']), None)
        if tot_row:
            qhdc_tot = tot_row['tong']
            qhdc_27 = tot_row['y2027']
            qhdc_28 = tot_row['y2028']
            strat_tot = tot_row['strat_tot']
            strat_26 = tot_row['strat_26']
            strat_27 = tot_row['strat_27']
            strat_28 = tot_row['strat_28']
            strat_29 = tot_row['strat_29']
            strat_30 = tot_row['strat_30']

            strat_27_28 = strat_27 + strat_28
            delta = qhdc_tot - strat_27_28
            rate = round((qhdc_tot / strat_27_28 * 100), 1) if strat_27_28 > 0 else 0.0

            status = 'Vừa khớp'
            status_color = 'green'
            if strat_27_28 > 0:
                if delta > 0.05:
                    status = f'Vượt CL +{delta:.2f} M$'
                    status_color = 'orange'
                elif delta < -0.05:
                    status = f'Dưới CL {delta:.2f} M$'
                    status_color = 'blue'

            comparison.append({
                'title': t['title'],
                'qhdc_2027': qhdc_27,
                'qhdc_2028': qhdc_28,
                'qhdc_total': qhdc_tot,
                'strat_2026': strat_26,
                'strat_2027': strat_27,
                'strat_2028': strat_28,
                'strat_2029': strat_29,
                'strat_2030': strat_30,
                'strat_27_28': strat_27_28,
                'strat_5y': strat_tot,
                'delta': delta,
                'rate': rate,
                'status': status,
                'status_color': status_color
            })

    return comparison

TH_DV_SECTORS = ['VT', 'ML', 'CDBR', 'CNTT', 'TD', 'IP', 'CD', 'HT']

TH_DV_ROW_DEFS = [
    {'id': 'R04', 'type': 'grand_total', 'stt': '*', 'nd': 'Tổng', 'madv': None},
    {'id': 'R05', 'type': 'share_header', 'stt': '', 'nd': 'Tỷ trọng theo mảng (%)', 'madv': None},
    {'id': 'R06', 'type': 'group', 'stt': 'I', 'nd': 'Công nghệ mới triển khai diện rộng cho kinh doanh', 'madv': None, 'children': ['R07', 'R08', 'R09', 'R10']},
    {'id': 'R07', 'type': 'leaf', 'stt': '1', 'nd': 'Mạng 5G', 'madv': '5G'},
    {'id': 'R08', 'type': 'leaf', 'stt': '2', 'nd': 'XGSPON', 'madv': 'XGSPON'},
    {'id': 'R09', 'type': 'leaf', 'stt': '3', 'nd': 'AI/GPU', 'madv': 'AI/GPU'},
    {'id': 'R10', 'type': 'leaf', 'stt': '4', 'nd': 'Private Mobile Network', 'madv': 'PMN'},
    {'id': 'R11', 'type': 'group', 'stt': 'II', 'nd': 'Mở rộng mạng lưới hiện tại cho kinh doanh', 'madv': None, 'children': ['R12', 'R15', 'R20', 'R21', 'R22']},
    {'id': 'R12', 'type': 'subgroup', 'stt': '1', 'nd': 'Mạng 2/3/4G', 'madv': None, 'children': ['R13', 'R14']},
    {'id': 'R13', 'type': 'leaf', 'stt': '-', 'nd': '2/3/4G vùng phủ', 'madv': 'VPDD'},
    {'id': 'R14', 'type': 'leaf', 'stt': '-', 'nd': '2/3/4G dung lượng', 'madv': 'DLDD'},
    {'id': 'R15', 'type': 'subgroup', 'stt': '2', 'nd': 'Mạng BRCĐ&TH', 'madv': None, 'children': ['R16', 'R17', 'R18', 'R19']},
    {'id': 'R16', 'type': 'leaf', 'stt': '-', 'nd': 'Vùng phủ GPON', 'madv': 'VPGPON'},
    {'id': 'R17', 'type': 'leaf', 'stt': '-', 'nd': 'Dung lượng Internet', 'madv': 'DLGPON'},
    {'id': 'R18', 'type': 'leaf', 'stt': '-', 'nd': 'Dung lượng Truyền hình', 'madv': 'TH'},
    {'id': 'R19', 'type': 'leaf', 'stt': '-', 'nd': 'Kênh truyền', 'madv': 'KT'},
    {'id': 'R20', 'type': 'leaf', 'stt': '3', 'nd': 'Phục vụ kinh doanh Cloud', 'madv': 'CLOUD_KD'},
    {'id': 'R21', 'type': 'leaf', 'stt': '4', 'nd': 'Data Center', 'madv': 'DC'},
    {'id': 'R22', 'type': 'subgroup', 'stt': '5', 'nd': 'Triển khai hạ tầng CNTT', 'madv': None, 'children': ['R23', 'R24', 'R25']},
    {'id': 'R23', 'type': 'leaf', 'stt': '-', 'nd': 'Trực tiếp kinh doanh', 'madv': 'TTKD'},
    {'id': 'R24', 'type': 'leaf', 'stt': '-', 'nd': 'Hỗ trợ kinh doanh và quản trị', 'madv': 'HTKD'},
    {'id': 'R25', 'type': 'leaf', 'stt': '-', 'nd': 'Phần mềm', 'madv': 'PM'},
    {'id': 'R26', 'type': 'group', 'stt': 'III', 'nd': 'Đảm bảo VHKT, kiên cố, bền vững trong thiên tai', 'madv': None, 'children': ['R27', 'R28', 'R29', 'R30']},
    {'id': 'R27', 'type': 'leaf', 'stt': '1', 'nd': 'Đảm bảo dự phòng ƯCTT', 'madv': 'UCTT'},
    {'id': 'R28', 'type': 'leaf', 'stt': '2', 'nd': 'Đảm bảo VHKT và nâng cao chất lượng mạng', 'madv': 'VHKT'},
    {'id': 'R29', 'type': 'leaf', 'stt': '3', 'nd': 'Củng cố kiên cố thường trình', 'madv': 'KCTT'},
    {'id': 'R30', 'type': 'leaf', 'stt': '4', 'nd': 'Bền vững mạng lưới trong thiên tai', 'madv': 'PCTT'},
    {'id': 'R31', 'type': 'leaf', 'stt': 'IV', 'nd': 'Ví điện tử', 'madv': 'VI'}
]

def compute_th_theo_dv(matrix_27, matrix_28):
    row_data = {}
    for r in TH_DV_ROW_DEFS:
        row_data[r['id']] = {
            'id': r['id'],
            'type': r['type'],
            'stt': r['stt'],
            'nd': r['nd'],
            'madv': r['madv'],
            'y2027': {s: 0.0 for s in TH_DV_SECTORS} | {'total': 0.0, 'share': 0.0},
            'y2028': {s: 0.0 for s in TH_DV_SECTORS} | {'total': 0.0, 'share': 0.0},
            'yTotal': {s: 0.0 for s in TH_DV_SECTORS} | {'total': 0.0, 'share': 0.0}
        }

    # 1. Fill leaf rows
    for r in TH_DV_ROW_DEFS:
        if r['madv']:
            d = row_data[r['id']]
            code = norm_madv(r['madv'])
            for s in TH_DV_SECTORS:
                s_key = 'CĐBR' if s == 'CDBR' else ('CĐ' if s == 'CD' else s)
                v27 = matrix_27.get((s_key, code), 0.0) or matrix_27.get((s, code), 0.0)
                v28 = matrix_28.get((s_key, code), 0.0) or matrix_28.get((s, code), 0.0)
                d['y2027'][s] = v27
                d['y2028'][s] = v28
                d['yTotal'][s] = v27 + v28
            d['y2027']['total'] = sum(d['y2027'][s] for s in TH_DV_SECTORS)
            d['y2028']['total'] = sum(d['y2028'][s] for s in TH_DV_SECTORS)
            d['yTotal']['total'] = sum(d['yTotal'][s] for s in TH_DV_SECTORS)

    # 2. Subgroups (R12, R15, R22)
    for sub_id in ['R12', 'R15', 'R22']:
        def_r = next(r for r in TH_DV_ROW_DEFS if r['id'] == sub_id)
        d = row_data[sub_id]
        for s in TH_DV_SECTORS:
            d['y2027'][s] = sum(row_data[c]['y2027'][s] for c in def_r['children'])
            d['y2028'][s] = sum(row_data[c]['y2028'][s] for c in def_r['children'])
            d['yTotal'][s] = sum(row_data[c]['yTotal'][s] for c in def_r['children'])
        d['y2027']['total'] = sum(d['y2027'][s] for s in TH_DV_SECTORS)
        d['y2028']['total'] = sum(d['y2028'][s] for s in TH_DV_SECTORS)
        d['yTotal']['total'] = sum(d['yTotal'][s] for s in TH_DV_SECTORS)

    # 3. Groups (R06, R11, R26)
    for grp_id in ['R06', 'R11', 'R26']:
        def_r = next(r for r in TH_DV_ROW_DEFS if r['id'] == grp_id)
        d = row_data[grp_id]
        for s in TH_DV_SECTORS:
            d['y2027'][s] = sum(row_data[c]['y2027'][s] for c in def_r['children'])
            d['y2028'][s] = sum(row_data[c]['y2028'][s] for c in def_r['children'])
            d['yTotal'][s] = sum(row_data[c]['yTotal'][s] for c in def_r['children'])
        d['y2027']['total'] = sum(d['y2027'][s] for s in TH_DV_SECTORS)
        d['y2028']['total'] = sum(d['y2028'][s] for s in TH_DV_SECTORS)
        d['yTotal']['total'] = sum(d['yTotal'][s] for s in TH_DV_SECTORS)

    # 4. Grand Total R04
    r04 = row_data['R04']
    for s in TH_DV_SECTORS:
        r04['y2027'][s] = sum(row_data[c]['y2027'][s] for c in ['R06', 'R11', 'R26', 'R31'])
        r04['y2028'][s] = sum(row_data[c]['y2028'][s] for c in ['R06', 'R11', 'R26', 'R31'])
        r04['yTotal'][s] = sum(row_data[c]['yTotal'][s] for c in ['R06', 'R11', 'R26', 'R31'])
    r04['y2027']['total'] = sum(r04['y2027'][s] for s in TH_DV_SECTORS)
    r04['y2028']['total'] = sum(r04['y2028'][s] for s in TH_DV_SECTORS)
    r04['yTotal']['total'] = sum(r04['yTotal'][s] for s in TH_DV_SECTORS)
    r04['y2027']['share'] = 100.0
    r04['y2028']['share'] = 100.0
    r04['yTotal']['share'] = 100.0

    # 5. Share row R05
    r05 = row_data['R05']
    for s in TH_DV_SECTORS:
        r05['y2027'][s] = round((r04['y2027'][s] / r04['y2027']['total'] * 100), 2) if r04['y2027']['total'] > 0 else 0.0
        r05['y2028'][s] = round((r04['y2028'][s] / r04['y2028']['total'] * 100), 2) if r04['y2028']['total'] > 0 else 0.0
        r05['yTotal'][s] = round((r04['yTotal'][s] / r04['yTotal']['total'] * 100), 2) if r04['yTotal']['total'] > 0 else 0.0
    r05['y2027']['total'] = 100.0
    r05['y2028']['total'] = 100.0
    r05['yTotal']['total'] = 100.0

    # 6. Row shares
    for r in TH_DV_ROW_DEFS:
        if r['id'] in ['R04', 'R05']:
            continue
        d = row_data[r['id']]
        d['y2027']['share'] = round((d['y2027']['total'] / r04['y2027']['total'] * 100), 2) if r04['y2027']['total'] > 0 else 0.0
        d['y2028']['share'] = round((d['y2028']['total'] / r04['y2028']['total'] * 100), 2) if r04['y2028']['total'] > 0 else 0.0
        d['yTotal']['share'] = round((d['yTotal']['total'] / r04['yTotal']['total'] * 100), 2) if r04['yTotal']['total'] > 0 else 0.0

    return [row_data[r['id']] for r in TH_DV_ROW_DEFS]
